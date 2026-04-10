from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


SRC_PATH = Path(r"Lab6/AnnealingAndAntReport.docx")
OUT_PATH = Path(r"Lab6/AnnealingAndAntReport_updated.docx")
RESEARCH_INDEX = 103
CONCLUSION_INDEX = 178


CONTENT = [
    ("text", "Метод отжига"),
    (
        "text",
        "На малом графе figure1 обе модификации отжига показали одинаковое среднее значение "
        "лучшего маршрута 16.5. Подбор параметров здесь не изменял качество решения, а только "
        "убирал избыточные вычисления: геометрический вариант сократился по среднему времени с "
        "228.3 до 3.65 мс, а вариант Коши - с 1935.6 до 3.82 мс. Это означает, что для "
        "небольшого контрольного графа влияние параметров проявляется прежде всего во времени "
        "выполнения, а не в качестве найденного пути.",
    ),
    ("text", "Влияние модификации:"),
    (
        "text",
        "На графе berlin52 геометрическое охлаждение оказалось сильнее: после настройки средняя "
        "длина лучшего маршрута уменьшилась с 2354.0 до 2191.2, тогда как модификация Коши дала "
        "2199.2. На world666 ситуация обратная: именно отжиг Коши показал лучший результат "
        "353582.0 против 366435.7 у настроенного геометрического варианта. Следовательно, "
        "геометрическое охлаждение лучше работает на среднем графе, а модификация Коши "
        "оказывается полезнее на большом графе, где требуется более длительное исследование "
        "пространства решений.",
    ),
    (
        "image",
        r"c:\Users\kameel\Repositories\Algorithms-and-data-structures\Lab6\benchmark_results\sensitivity_final_v2_20260326\sa\berlin52\baseline_vs_tuned.png",
    ),
    (
        "caption",
        "Рис. 8 Отжиг на berlin52: влияние настройки параметров и сравнение модификаций",
    ),
    ("text", "Влияние коэффициента alpha и числа итераций на температуру:"),
    (
        "text",
        "Для геометрического отжига на berlin52 наилучшее качество дали более медленное "
        "охлаждение и увеличенное число переходов на каждой температуре. По серии запусков "
        "лучшей оказалась конфигурация alpha = 0.975 и iterations_per_temperature = 150; она "
        "даёт заметно меньшую среднюю длину маршрута, чем базовые значения alpha = 0.95 и "
        "iterations_per_temperature = 50. Параметр min_temperature также сместился к 0.2, что "
        "говорит о полезности остановки раньше зоны почти нулевых изменений.",
    ),
    (
        "image",
        r"c:\Users\kameel\Repositories\Algorithms-and-data-structures\Lab6\benchmark_results\sensitivity_final_v2_20260326\sa\berlin52\sa_geometric\plots\alpha.png",
    ),
    (
        "caption",
        "Рис. 9 Влияние коэффициента alpha в геометрическом отжиге на графе berlin52",
    ),
    ("text", "Влияние начальной температуры в модификации Коши:"),
    (
        "text",
        "На world666 модификация Коши чувствительна прежде всего к initial_temperature. При "
        "росте начальной температуры от 100 до 3000 средняя длина лучшего маршрута устойчиво "
        "уменьшается, и лучший результат достигается при initial_temperature = 3000. Однако "
        "улучшение качества оплачивается временем: tuned-конфигурация даёт выигрыш 2.84% "
        "относительно baseline, но увеличивает среднее время с 2.69 до 66.78 с.",
    ),
    (
        "image",
        r"c:\Users\kameel\Repositories\Algorithms-and-data-structures\Lab6\benchmark_results\sensitivity_final_v2_20260326\sa\world666\sa_cauchy\plots\initial_temperature.png",
    ),
    (
        "caption",
        "Рис. 10 Влияние initial_temperature в модификации отжига Коши на графе world666",
    ),
    ("text", "Сравнение настроенных вариантов отжига:"),
    (
        "text",
        "После настройки геометрический отжиг остаётся более экономичным по времени, а отжиг "
        "Коши даёт лучшее качество на большом графе. На world666 настроенный вариант Коши лучше "
        "настроенного геометрического на 3.51%, но работает примерно в 15.3 раза дольше. "
        "Поэтому модификацию Коши разумно использовать тогда, когда качество маршрута важнее "
        "времени выполнения.",
    ),
    (
        "image",
        r"c:\Users\kameel\Repositories\Algorithms-and-data-structures\Lab6\benchmark_results\sensitivity_final_v2_20260326\sa\world666\tuned_variant_comparison.png",
    ),
    ("caption", "Рис. 11 Сравнение настроенных вариантов отжига на графе world666"),
    ("text", "Муравьиный алгоритм"),
    (
        "text",
        "На контрольном графе figure1 оба варианта муравьиного алгоритма стабильно находят "
        "оптимальный цикл длины 14.0. Поэтому на малом графе модификация \"элитные муравьи\" "
        "не улучшает качество решения: различие проявляется только во времени, причём базовый "
        "вариант оказался быстрее (5.01 мс против 7.91 мс).",
    ),
    ("text", "Влияние модификации:"),
    (
        "text",
        "На berlin52 элитный вариант лучше базового как до, так и после настройки. После "
        "подбора параметров элитный алгоритм дал среднюю длину 1887.0, а базовый - 1922.4. На "
        "world666 преимущество элитной модификации стало ещё заметнее: 580918.0 против 619499.0 "
        "у обычного алгоритма. Таким образом, дополнительное усиление лучшего маршрута "
        "феромоном оказывается полезным на средних и больших графах.",
    ),
    (
        "image",
        r"c:\Users\kameel\Repositories\Algorithms-and-data-structures\Lab6\benchmark_results\sensitivity_final_v2_20260326\aco_fb\berlin52\baseline_vs_tuned.png",
    ),
    (
        "caption",
        "Рис. 12 Муравьиный алгоритм на berlin52: сравнение базового и элитного вариантов до и после настройки",
    ),
    ("text", "Влияние числа итераций и размера колонии:"),
    (
        "text",
        "Для базового муравьиного алгоритма на berlin52 наиболее заметный выигрыш даёт "
        "увеличение числа итераций: при переходе от 20 к 150 средняя длина маршрута устойчиво "
        "уменьшается. При этом рост ant_count помогает только до умеренных значений; лучшим "
        "оказался диапазон около 30 муравьёв, тогда как дальнейшее увеличение колонии не даёт "
        "сопоставимого улучшения качества.",
    ),
    (
        "image",
        r"c:\Users\kameel\Repositories\Algorithms-and-data-structures\Lab6\benchmark_results\sensitivity_final_v2_20260326\aco_fb\berlin52\ant_basic\plots\iterations.png",
    ),
    (
        "caption",
        "Рис. 13 Влияние числа итераций в базовом муравьином алгоритме на графе berlin52",
    ),
    ("text", "Влияние числа элитных муравьёв:"),
    (
        "text",
        "На berlin52 элитный алгоритм лучше всего работает при 5-8 элитных муравьях, а лучший "
        "результат получен при elite_ants = 8. Рост этого параметра уменьшает среднюю длину "
        "маршрута по сравнению с режимом без выраженного элитизма, но сам выигрыш умеренный, "
        "потому что baseline-конфигурация уже находилась близко к хорошей области параметров.",
    ),
    (
        "image",
        r"c:\Users\kameel\Repositories\Algorithms-and-data-structures\Lab6\benchmark_results\sensitivity_final_v2_20260326\aco_fb\berlin52\ant_elitist\plots\elite_ants.png",
    ),
    (
        "caption",
        "Рис. 14 Влияние числа элитных муравьёв в модифицированном алгоритме на графе berlin52",
    ),
    ("text", "Влияние параметров на большом графе:"),
    (
        "text",
        "На world666 муравьиный алгоритм оказался наиболее чувствителен к параметрам, "
        "управляющим балансом между запоминанием и обновлением следов. Лучшая конфигурация для "
        "обычного варианта использует ant_count = 16, iterations = 12, distance_importance = "
        "3.0, evaporation_intensity = 4.0 и пониженное pheromone_importance = 0.5. Для "
        "элитного варианта лучшими оказались те же ant_count = 16 и iterations = 12, но при "
        "pheromone_importance = 1.5 и elite_ants = 5. Это дало крупный выигрыш относительно "
        "baseline: 49.46% для обычного варианта и 52.58% для элитного.",
    ),
    (
        "image",
        r"c:\Users\kameel\Repositories\Algorithms-and-data-structures\Lab6\benchmark_results\sensitivity_final_v2_20260326\aco_world\world666\baseline_vs_tuned.png",
    ),
    (
        "caption",
        "Рис. 15 Муравьиный алгоритм на world666: сравнение baseline и tuned-конфигураций",
    ),
    ("text", "Вывод по исследованию - оптимальный алгоритм"),
    (
        "text",
        "После отдельной настройки универсального победителя для всех графов не получилось. На "
        "figure1 лучшие результаты дали оба муравьиных алгоритма: они находят оптимальный цикл "
        "длины 14.0, тогда как отжиг остаётся на 16.5. На berlin52 лидером стал алгоритм "
        "\"элитных\" муравьёв с результатом 1887.0; это на 13.88% лучше лучшего варианта "
        "отжига, но примерно в 5 раз медленнее геометрического отжига. На world666 ситуация "
        "обратная: даже быстрый геометрический отжиг лучше элитного муравьиного алгоритма на "
        "36.92% и при этом работает примерно в 1.9 раза быстрее, а наилучшее качество вообще "
        "даёт отжиг Коши - 353582.0. Поэтому на малых и средних графах предпочтительнее "
        "муравьиный алгоритм, особенно его элитная модификация, а на большом графе более "
        "удачным оказался метод имитации отжига.",
    ),
]


def add_text(anchor_paragraph, text, style_names, style, align=None):
    paragraph = anchor_paragraph.insert_paragraph_before()
    if style in style_names:
        paragraph.style = style
    if text:
        paragraph.add_run(text)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def add_image(anchor_paragraph, image_path):
    paragraph = anchor_paragraph.insert_paragraph_before()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(image_path, width=Inches(5.75))
    return paragraph


def main():
    doc = Document(SRC_PATH)
    style_names = {style.name for style in doc.styles}
    caption_style = "Normal (Web)" if "Normal (Web)" in style_names else "Normal"
    normal_style = "Normal" if "Normal" in style_names else None

    paragraphs = doc.paragraphs
    research_para = paragraphs[RESEARCH_INDEX]
    conclusion_heading = paragraphs[CONCLUSION_INDEX]

    body = research_para._element.getparent()
    remove_mode = False
    for child in list(body):
        if child == research_para._element:
            remove_mode = True
            continue
        if child == conclusion_heading._element:
            break
        if remove_mode:
            body.remove(child)

    for item_type, payload in CONTENT:
        if item_type == "text":
            add_text(conclusion_heading, payload, style_names, normal_style or "Normal")
        elif item_type == "caption":
            add_text(
                conclusion_heading,
                payload,
                style_names,
                caption_style,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
        elif item_type == "image":
            add_image(conclusion_heading, payload)
        else:
            raise RuntimeError(f"Unknown content item type: {item_type}")

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
