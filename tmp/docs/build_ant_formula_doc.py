from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
TMP_DIR = ROOT / "tmp" / "docs" / "ant_formula_assets"
OUTPUT_DIR = ROOT / "output" / "doc"
DOCX_PATH = OUTPUT_DIR / "lab6_ant_formulas_explained.docx"


FORMULAS = [
    {
        "number": "(1)",
        "latex": r"\eta_{ij}=\frac{1}{d_{ij}}",
        "notes": [
            "ηᵢⱼ — оценка привлекательности ребра (i, j) по расстоянию;",
            "dᵢⱼ — расстояние между вершинами i и j.",
        ],
    },
    {
        "number": "(2)",
        "latex": r"\eta_{ij}^{(\beta)}=\left(\frac{1}{d_{ij}}\right)^\beta",
        "notes": [
            "ηᵢⱼ⁽ᵝ⁾ — оценка расстояния с учётом параметра β;",
            "β — коэффициент влияния расстояния.",
        ],
    },
    {
        "number": "(3)",
        "latex": r"a_{ij}=\tau_{ij}^{\alpha}\eta_{ij}^{(\beta)}",
        "notes": [
            "aᵢⱼ — общая привлекательность ребра (i, j);",
            "τᵢⱼ — количество феромона на ребре (i, j);",
            "α — коэффициент влияния феромона.",
        ],
    },
    {
        "number": "(4)",
        "latex": r"p_{ij}=\frac{a_{ij}}{\sum_{l\in C} a_{il}}",
        "notes": [
            "pᵢⱼ — вероятность перехода из вершины i в вершину j;",
            "l — индекс возможной следующей вершины;",
            "C — множество допустимых вершин для перехода.",
        ],
    },
    {
        "number": "(5)",
        "latex": r"L_k=\sum_i d(\pi_i,\pi_{i+1})",
        "notes": [
            "Lₖ — длина маршрута, построенного k-м муравьём;",
            "π — маршрут муравья;",
            "πᵢ — вершина на позиции i в маршруте;",
            "Σᵢ — суммирование по всем позициям маршрута;",
            "k — номер муравья.",
        ],
    },
    {
        "number": "(6)",
        "latex": r"\rho=\frac{\mathrm{evaporation\_intensity}}{10}",
        "notes": [
            "ρ — коэффициент испарения феромона;",
            "evaporation_intensity — параметр интенсивности испарения.",
        ],
    },
    {
        "number": "(7)",
        "latex": r"\tau_{ij}\leftarrow(1-\rho)\tau_{ij}",
        "notes": [
            "← — операция обновления значения феромона.",
        ],
    },
    {
        "number": "(8)",
        "latex": r"\Delta\tau_{ij}^{(k)}=\frac{Q}{L_k}",
        "notes": [
            "Δτᵢⱼ⁽ᵏ⁾ — количество феромона, добавляемое k-м муравьём на ребро (i, j);",
            "Q — коэффициент добавления феромона.",
        ],
    },
    {
        "number": "(9)",
        "latex": r"\tau_{ij}\leftarrow\tau_{ij}+\sum_k \Delta\tau_{ij}^{(k)}",
        "notes": [
            "Σₖ — суммирование вкладов всех муравьёв.",
        ],
    },
    {
        "number": "(10)",
        "latex": r"\Delta\tau_{ij}^{\mathrm{best}}=\frac{eQ}{L_{\mathrm{best}}}",
        "notes": [
            "Δτᵢⱼᵇᵉˢᵗ — дополнительный феромон для лучшего маршрута;",
            "Lᵇᵉˢᵗ — длина лучшего найденного маршрута;",
            "e — число элитных муравьёв.",
        ],
    },
]


def render_formula(latex: str, output_path: Path) -> None:
    figure = plt.figure(figsize=(0.01, 0.01))
    text = figure.text(0, 0, f"${latex}$", fontsize=26)
    figure.canvas.draw()
    bbox = text.get_window_extent(renderer=figure.canvas.get_renderer())
    expanded = bbox.expanded(1.08, 1.18)
    figure.savefig(
        output_path,
        dpi=300,
        transparent=True,
        bbox_inches=expanded.transformed(figure.dpi_scale_trans.inverted()),
        pad_inches=0.02,
    )
    plt.close(figure)


def set_cell_borderless(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_table_borderless(table) -> None:
    table_pr = table._tbl.tblPr
    borders = table_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def configure_page(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)


def add_formula_row(table, image_path: Path, number: str, notes: list[str]) -> None:
    row = table.add_row()
    left_cell = row.cells[0]
    right_cell = row.cells[1]
    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_borderless(left_cell)
    set_cell_borderless(right_cell)

    formula_paragraph = left_cell.paragraphs[0]
    formula_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula_paragraph.paragraph_format.space_before = Pt(0)
    formula_paragraph.paragraph_format.space_after = Pt(0)
    formula_paragraph.add_run().add_picture(str(image_path), width=Inches(3.2))

    for note in notes:
        paragraph = left_cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = paragraph.add_run(note)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(12)

    number_paragraph = right_cell.paragraphs[0]
    number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_paragraph.paragraph_format.space_before = Pt(6)
    number_paragraph.paragraph_format.space_after = Pt(0)
    run = number_paragraph.add_run(number)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(14)


def build_document() -> Path:
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    image_paths = []
    for index, item in enumerate(FORMULAS, start=1):
        image_path = TMP_DIR / f"formula_{index:02d}.png"
        render_formula(item["latex"], image_path)
        image_paths.append(image_path)

    document = Document()
    configure_page(document)

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)

    section = document.sections[0]
    available_width = section.page_width - section.left_margin - section.right_margin
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = int(available_width * 0.84)
    table.columns[1].width = int(available_width * 0.16)
    set_table_borderless(table)

    for image_path, item in zip(image_paths, FORMULAS, strict=True):
        add_formula_row(table, image_path, item["number"], item["notes"])

    document.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = build_document()
    print(path)
